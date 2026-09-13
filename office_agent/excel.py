from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
import hashlib
import json
import multiprocessing
import shutil
import uuid
import time

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from .config import AppError
from .intent import valid_period

HEADERS = ["OrderID", "Date", "Customer", "Region", "Product", "Quantity", "UnitPrice", "Discount", "Status", "Revenue"]
REGIONS = ("Bắc", "Trung", "Nam")
DETAIL_HEADERS = ["OrderID", "Date", "Region", "Quantity", "UnitPrice", "Discount", "Revenue"]


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def money(value):
    try:
        if isinstance(value, bool):
            raise ValueError
        d = Decimal(str(value))
        if not d.is_finite():
            raise ValueError
        return d
    except (InvalidOperation, ValueError):
        raise AppError("Nguồn có số không hợp lệ.") from None


def inspect_source(path):
    try:
        book = load_workbook(path, data_only=False, keep_links=True)
    except Exception:
        raise AppError("Không mở được source.xlsx; kiểm tra file hoặc mật khẩu.") from None
    try:
        if book._external_links or "Orders" not in book.sheetnames:
            raise AppError("Nguồn cần sheet Orders và không có external links.")
        ws = book["Orders"]
        if ws.max_row > 10001 or ws.max_column != 10:
            raise AppError("MVP nhận tối đa 10.000 dòng và đúng 10 cột Orders.")
        if [c.value for c in ws[1]] != HEADERS:
            raise AppError("Các cột Orders không đúng bộ dữ liệu đã chốt.")
        if "Orders" not in ws.tables or ws.tables["Orders"].ref != f"A1:J{ws.max_row}":
            raise AppError("Table Orders phải bao phủ toàn bộ A1:J của sheet.")
        for sheet in book:
            if sheet.max_row > 10001 or sheet.max_column > 100:
                raise AppError("Nguồn có sheet phụ vượt giới hạn MVP.")
            for row in sheet:
                for cell in row:
                    if cell.data_type == "f" and not (sheet.title == "Orders" and cell.column == 10 and cell.row >= 2):
                        raise AppError("Nguồn có công thức ngoài cột Revenue được phép.")
        ids, rows = set(), []
        for index, values in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            oid, when, customer, region, product, qty, price, discount, status, revenue = values
            if not isinstance(oid, str) or not oid.strip() or oid in ids:
                raise AppError(f"OrderID trống/trùng tại dòng {index}.")
            ids.add(oid)
            if not isinstance(when, (date, datetime)) or region not in REGIONS or status not in ("Completed", "Pending", "Cancelled"):
                raise AppError(f"Ngày/vùng/trạng thái không hợp lệ tại dòng {index}.")
            q, p, d = money(qty), money(price), money(discount)
            if q < 0 or q != q.to_integral_value() or p < 0 or not 0 <= d <= 1:
                raise AppError(f"Số lượng/giá/chiết khấu ngoài giới hạn tại dòng {index}.")
            # Do not evaluate arbitrary workbook formulas/macros. Only this reviewed formula.
            expected_formula = f"=ROUND(F{index}*G{index}*(1-H{index}),0)"
            if isinstance(revenue, str) and revenue.startswith("="):
                if revenue.upper().replace(" ", "") != expected_formula:
                    raise AppError(f"Công thức Revenue ngoài phạm vi tại dòng {index}.")
            expected = int((q * p * (1 - d)).quantize(Decimal(1), rounding=ROUND_HALF_UP))
            if not (isinstance(revenue, str) and revenue.startswith("=")) and money(revenue) != expected:
                raise AppError(f"Revenue không khớp phép tính tại dòng {index}.")
            # All other cells must be data, not formulas (including customer/product text).
            if any(isinstance(v, str) and v.startswith("=") for v in values[:9]):
                raise AppError(f"Công thức ngoài cột Revenue tại dòng {index}.")
            rows.append({"id": oid, "date": when.strftime("%Y-%m-%d"), "region": region,
                         "quantity": int(q), "unit_price": str(p), "discount": str(d),
                         "status": status, "revenue": expected})
        return rows
    finally:
        book.close()


def _excel_worker(path, connection, cancel):
    if cancel.is_set():
        connection.close()
        return
    import pythoncom
    import win32com.client
    import win32process
    pythoncom.CoInitialize()
    app = book = None
    try:
        if cancel.is_set():
            return
        app = win32com.client.DispatchEx("Excel.Application")
        connection.send(("pid", win32process.GetWindowThreadProcessId(app.Hwnd)[1]))
        if cancel.is_set():
            return
        app.Visible = False
        app.DisplayAlerts = False
        app.AutomationSecurity = 3
        book = app.Workbooks.Open(str(path.resolve()), UpdateLinks=0, ReadOnly=False)
        app.CalculateFullRebuild()
        book.Save()
        connection.send(("ok", None))
    except Exception:
        connection.send(("error", None))
    finally:
        if book is not None:
            try:
                book.Close(False)
            except Exception:
                pass
        if app is not None:
            try:
                app.Quit()
            except Exception:
                pass
        book = app = None
        pythoncom.CoUninitialize()
        connection.close()


def recalculate_snapshot(path, timeout=45):
    """Bound COM time; retain a handle only to the Excel instance we created."""
    import win32api
    import win32event
    ctx = multiprocessing.get_context("spawn")
    receiver, sender = ctx.Pipe(duplex=False)
    cancel = ctx.Event()
    worker = ctx.Process(target=_excel_worker, args=(path, sender, cancel), daemon=False)
    handle = None
    success = False
    deadline = time.monotonic() + timeout
    worker.start()
    sender.close()
    try:
        while time.monotonic() < deadline:
            if receiver.poll(0.1):
                try:
                    kind, value = receiver.recv()
                except EOFError:
                    break
                if kind == "pid":
                    handle = win32api.OpenProcess(0x0001 | 0x00100000, False, value)
                elif kind == "ok":
                    success = True
                elif kind == "error":
                    break
            if not worker.is_alive():
                break
        if worker.is_alive() and time.monotonic() >= deadline:
            cancel.set()
            # Let startup see cancellation or report its owned Excel handle before
            # terminating. Killing mid-DispatchEx can otherwise orphan Office.
            grace = time.monotonic() + 3
            while worker.is_alive() and time.monotonic() < grace:
                if receiver.poll(0.1):
                    try:
                        kind, value = receiver.recv()
                    except EOFError:
                        break
                    if kind == "pid" and handle is None:
                        handle = win32api.OpenProcess(0x0001 | 0x00100000, False, value)
            success = False
        worker.join(timeout=0.5)
        if worker.is_alive():
            success = False
        if not success:
            raise AppError("Excel không hoàn tất trong thời gian cho phép. Bản gốc được giữ nguyên; hãy thử lại.")
    finally:
        if worker.is_alive():
            worker.terminate()
        worker.join(timeout=3)
        worker.close()
        receiver.close()
        if handle is not None:
            try:
                if win32event.WaitForSingleObject(handle, 2000) == win32event.WAIT_TIMEOUT:
                    win32api.TerminateProcess(handle, 1)
                    win32event.WaitForSingleObject(handle, 2000)
            finally:
                handle.Close()


@dataclass(frozen=True)
class Artifact:
    path: Path
    digest: str
    source_digest: str
    snapshot: Path
    period: str
    count: int
    total: int
    warning: str


def format_sheet(ws):
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="176B78")
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(60, max(18, max(len(str(c.value or "")) for c in col) + 2))


def build_report(source, output_dir, period, run_id=None, *, excel_recalc=True):
    valid_period(period)
    run_id = run_id or uuid.uuid4().hex
    if not run_id.isalnum():
        raise AppError("Run ID không hợp lệ.")
    output_dir.mkdir(parents=True, exist_ok=True)
    run_dir = output_dir / run_id
    run_dir.mkdir(exist_ok=False)
    source_hash = sha(source)
    snapshot = run_dir / "source_snapshot.xlsx"
    shutil.copyfile(source, snapshot)
    if sha(snapshot) != source_hash or sha(source) != source_hash:
        raise AppError("Nguồn thay đổi trong lúc chụp bản sao; thử lại.")
    rows = inspect_source(snapshot)
    if excel_recalc:
        recalculate_snapshot(snapshot)
        cached = load_workbook(snapshot, data_only=True)
        try:
            if any(cached["Orders"].cell(i, 10).value != r["revenue"] for i, r in enumerate(rows, 2)):
                raise AppError("Kết quả Excel tính lại không khớp phép tính độc lập.")
        finally:
            cached.close()
    selected = [r for r in rows if r["date"].startswith(period) and r["status"] == "Completed"]
    if not selected:
        raise AppError("Kỳ này không có đơn Completed. Chưa tạo báo cáo; hãy chọn kỳ khác.")
    period_rows = [r for r in rows if r["date"].startswith(period)]
    warning = "Dữ liệu nguồn đến " + max(r["date"] for r in period_rows) + "; không khẳng định đã đủ toàn tháng."
    book = Workbook()
    summary = book.active
    summary.title = "Summary"
    summary.append(["Region", "CompletedOrders", "RevenueVND"])
    for region in REGIONS:
        group = [r for r in selected if r["region"] == region]
        summary.append([region, len(group), sum(r["revenue"] for r in group)])
    count, total = len(selected), sum(r["revenue"] for r in selected)
    summary.append(["TOTAL", count, total])
    details = book.create_sheet("Details")
    details.append(DETAIL_HEADERS)
    for r in selected:
        details.append([r["id"], datetime.fromisoformat(r["date"]), r["region"], r["quantity"], float(r["unit_price"]), float(r["discount"]), r["revenue"]])
        details.cell(details.max_row, 2).number_format = "yyyy-mm-dd"
        # Explicit text type prevents formula interpretation in external identifiers.
        details.cell(details.max_row, 1).data_type = "s"
    metadata = book.create_sheet("Metadata")
    metadata.append(["Field", "Value"])
    for row in [("Period", period), ("SourceSHA256", source_hash), ("RunID", run_id),
                ("CreatedAt", datetime.now().astimezone().isoformat()), ("Coverage", warning)]:
        metadata.append(row)
    for ws in book:
        format_sheet(ws)
    for row in summary.iter_rows(min_row=2, min_col=3, max_col=3):
        row[0].number_format = '#,##0" VND"'
    path = run_dir / f"bao_cao_doanh_thu_{period}.xlsx"
    staging = run_dir / "report_staging.xlsx"
    book.save(staging)
    book.close()
    # Independent verifier re-reads snapshot and persisted workbook.
    evidence = verify_report(snapshot, staging, period, source_hash)
    if sha(source) != source_hash:
        raise AppError("Nguồn đã đổi trong lúc tạo báo cáo. Chưa công bố artifact.")
    staging.rename(path)
    evidence.update({"report_sha256": sha(path), "snapshot_sha256": sha(snapshot), "excel_recalculated": excel_recalc})
    (run_dir / "verification.json").write_text(json.dumps(evidence, ensure_ascii=False, indent=2), encoding="utf-8")
    return Artifact(path, sha(path), source_hash, snapshot, period, count, total, warning)


def verify_report(snapshot, report, period, source_hash):
    """Recompute from source primitive columns, not producer output/totals or AI."""
    records = inspect_source(snapshot)
    expected = {}
    sums = {region: [0, 0] for region in REGIONS}
    for r in records:
        if r["date"][:7] != period or r["status"] != "Completed":
            continue
        value = int((Decimal(r["quantity"]) * Decimal(r["unit_price"]) * (1 - Decimal(r["discount"]))).quantize(Decimal(1), rounding=ROUND_HALF_UP))
        expected[r["id"]] = [r["date"], r["region"], r["quantity"], Decimal(r["unit_price"]), Decimal(r["discount"]), value]
        sums[r["region"]][0] += 1
        sums[r["region"]][1] += value
    book = load_workbook(report, data_only=False)
    try:
        if book.sheetnames != ["Summary", "Details", "Metadata"]:
            raise ValueError
        if any(cell.data_type == "f" for ws in book for row in ws for cell in row):
            raise ValueError
        summary = list(book["Summary"].values)
        wanted = [("Region", "CompletedOrders", "RevenueVND")]
        wanted += [(region, *sums[region]) for region in REGIONS]
        wanted += [("TOTAL", sum(v[0] for v in sums.values()), sum(v[1] for v in sums.values()))]
        if summary != wanted or list(next(book["Details"].values)) != DETAIL_HEADERS:
            raise ValueError
        seen = set()
        for oid, when, region, qty, price, discount, revenue in list(book["Details"].values)[1:]:
            actual = [when.strftime("%Y-%m-%d"), region, qty, money(price), money(discount), revenue]
            if oid in seen or expected.get(oid) != actual:
                raise ValueError
            seen.add(oid)
        if seen != set(expected):
            raise ValueError
        metadata = dict(list(book["Metadata"].values)[1:])
        if metadata.get("Period") != period or metadata.get("SourceSHA256") != source_hash:
            raise ValueError
    except (ValueError, KeyError, TypeError, AttributeError):
        raise AppError("Kiểm chứng thất bại: báo cáo không khớp nguồn/chi tiết/tổng.") from None
    finally:
        book.close()
    return {"verified": True, "period": period, "source_sha256": source_hash,
            "count": len(expected), "total": sum(v[-1] for v in expected.values())}
