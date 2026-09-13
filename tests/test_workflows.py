from dataclasses import replace
from datetime import date
import time

import pytest
from openpyxl import load_workbook

from office_agent.config import AppError, Settings
from office_agent.engine import Engine
from office_agent.excel import build_report, inspect_source, sha, verify_report
from office_agent.intent import command, local_intent, validate_candidate
from office_agent.mail import contacts, save_draft


class NoNetwork:
    def parse(self, *args):
        raise AssertionError("Deterministic workflow must not call API")


def test_report_rounding_filter_and_original_unchanged(settings, artifact):
    assert artifact.count == 2
    assert artifact.total == 2051  # Excel ROUND: 250.5 rounds to 251.
    assert sha(settings.inputs / "source.xlsx") == artifact.source_digest
    assert verify_report(artifact.snapshot, artifact.path, "2026-08", artifact.source_digest)["verified"]


@pytest.mark.parametrize("sheet,cell,value", [("Summary", "C5", 9), ("Details", "G2", 9), ("Details", "A3", "A"), ("Metadata", "B2", "2026-07")])
def test_verifier_rejects_persisted_corruption(artifact, sheet, cell, value):
    book = load_workbook(artifact.path)
    book[sheet][cell] = value
    book.save(artifact.path)
    book.close()
    with pytest.raises(AppError, match="Kiểm chứng"):
        verify_report(artifact.snapshot, artifact.path, artifact.period, artifact.source_digest)


@pytest.mark.parametrize("cell,value", [("J2", "=SUM(A1:A99)"), ("A3", "A"), ("H2", 2), ("C2", '=HYPERLINK("https://invalid")')])
def test_bad_source_is_rejected(settings, cell, value):
    source = settings.inputs / "source.xlsx"
    book = load_workbook(source)
    book["Orders"][cell] = value
    book.save(source)
    book.close()
    with pytest.raises(AppError):
        inspect_source(source)


def test_formulas_on_hidden_sheet_are_rejected(settings):
    source = settings.inputs / "source.xlsx"
    book = load_workbook(source)
    sheet = book.create_sheet("Hidden")
    sheet.sheet_state = "hidden"
    sheet["A1"] = '=WEBSERVICE("https://invalid")'
    book.save(source)
    book.close()
    with pytest.raises(AppError, match="công thức ngoài"):
        inspect_source(source)


def test_unsupported_filter_is_not_silently_ignored():
    assert local_intent("Tạo báo cáo tháng 8 năm 2026 tính Pending", []).workflow == "unsupported"


def test_no_data_does_not_produce_completed_report(settings):
    with pytest.raises(AppError, match="không có đơn"):
        build_report(settings.inputs / "source.xlsx", settings.outputs, "2026-01", excel_recalc=False)
    assert not list(settings.outputs.rglob("bao_cao*.xlsx"))


@pytest.mark.parametrize("text,expected", [("không gửi bản 1234", "no_send"), ("chưa xác nhận gửi", "no_send"), ("Gửi bản 1234", "confirm:1234"), ("xác nhận gửi", "challenge"), ("Dừng gửi, chỉ lưu nháp", "draft"), ("Trợ lý, dừng lại!", "stop")])
def test_controls_require_exact_confirmation(text, expected):
    assert command(text) == expected


def test_spoken_alias_and_ambiguity(settings):
    aliases = contacts(settings.inputs / "contacts.csv")
    assert local_intent("Gửi cho Lan kế toán", aliases).recipient == "lan_ketoan"
    assert local_intent("Lan kế toán", aliases).recipient == "lan_ketoan"
    unclear = local_intent("Gửi cho Lan", aliases)
    assert unclear.recipient is None and unclear.question
    assert local_intent("tháng trước", aliases, date(2026, 1, 1)).period == "2025-12"
    assert local_intent("đổi tháng 8", aliases).question


def test_multiturn_email_and_source_refresh(settings, journal):
    e = Engine(settings, journal, NoNetwork(), excel_recalc=False)
    assert "tháng nào" in e.handle("Gửi báo cáo cho Lan kế toán")
    e.handle("tháng 8 năm 2026")
    assert e.preview.contact.alias == "lan_ketoan"
    assert e.preview.attachment.exists()
    first = e.artifact
    old_code = e.gate.approval.code
    e.handle("Gửi cho Lan")
    assert e.preview is None and e.gate.approval is None
    e.handle("Lan kế toán")
    assert e.preview is not None
    e.handle("Chỉ lưu nháp")
    assert list(first.path.parent.glob("draft_*.eml"))
    assert e.gate.approval is None
    source = settings.inputs / "source.xlsx"
    book = load_workbook(source)
    book["Orders"]["F2"] = 3
    book.save(source)
    book.close()
    e.handle("Tạo báo cáo tháng 8 năm 2026")
    assert e.artifact.total == 2951
    assert e.artifact.path != first.path
    assert e.preview is None


def test_stopped_engine_and_old_voice_confirmation(settings, journal):
    e = Engine(settings, journal, NoNetwork(), excel_recalc=False)
    e.handle("Gửi báo cáo tháng 8 năm 2026 cho Lan kế toán")
    code = e.gate.approval.code
    with pytest.raises(AppError, match="thu trước"):
        e.handle("Gửi bản " + code, captured_at=time.monotonic() - 300)
    e.handle("Dừng lại")
    with pytest.raises(AppError, match="tạm dừng"):
        e.handle("Tạo báo cáo tháng 7 năm 2026")
    e.handle("Tiếp tục")
    assert e.gate.approval is not None


def test_draft_has_attachment_and_no_network(preview):
    from email.parser import BytesParser
    from email.policy import default
    message = BytesParser(policy=default).parsebytes(save_draft(preview).read_bytes())
    assert message["To"] == "lan@demo.test"
    assert len(list(message.iter_attachments())) == 1


@pytest.mark.parametrize("data", [
    {"workflow": "send_now", "period": None, "recipient": None},
    {"workflow": "email", "period": None, "recipient": "outsider"},
    {"workflow": "email", "period": None, "recipient": None, "approved": True},
    {"workflow": "excel_report", "period": "2026-13", "recipient": None},
])
def test_ai_cannot_expand_scope(data):
    with pytest.raises(AppError):
        validate_candidate(data, ["lan_ketoan"])


def test_configuration_blocks_paid_and_external_paths(tmp_path):
    (tmp_path / ".env").write_text("OPENROUTER_TEXT_MODEL=paid-model\n")
    with pytest.raises(AppError, match=":free"):
        Settings.load(tmp_path)
    (tmp_path / ".env").write_text("INPUTS_DIR=../outside\n")
    with pytest.raises(AppError, match="ngoài"):
        Settings.load(tmp_path)
